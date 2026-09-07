import asyncio
import io
import tempfile
import sqlite3
import unittest
from pathlib import Path
from fastapi import HTTPException
from openpyxl import load_workbook
from app import database, main, tenancies as api
from app.schemas import (PropertyCreate, PropertyUpdate, PaymentCreate, TenantCreate,
                         TenancyCreate, RentRateCreate, DepositRefundCreate)


class TenancyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_path = database.DB_PATH
        database.DB_PATH = Path(self.tmp.name)/'test.db'
        database.init_db()
        self.prop = main.create_property(PropertyCreate(name='House 1',monthly_rent=12000,tenant_name='Legacy person'))
        self.ravi = api.create_tenant(TenantCreate(name='Ravi',phone='123'))
        self.priya = api.create_tenant(TenantCreate(name='Priya'))

    def tearDown(self):
        database.DB_PATH = self.old_path
        self.tmp.cleanup()

    def stay(self, tenant, start, end=None, rent=10000):
        return api.create_tenancy(TenancyCreate(property_id=self.prop['id'],tenant_id=tenant['id'],start_date=start,end_date=end,initial_rent=rent,deposit=20000))

    def pay(self, stay, month, amount):
        return main.create_payment(PaymentCreate(property_id=self.prop['id'],tenancy_id=stay['id'],rental_month=month,amount=amount,paid_on='2026-09-01'))

    def test_old_tenants_and_rates_drive_historical_balances(self):
        ravi = self.stay(self.ravi,'2024-01-01','2024-12-31')
        rate = api.create_rate(RentRateCreate(tenancy_id=ravi['id'],effective_month='2024-07',amount=11000))
        priya = self.stay(self.priya,'2025-01-01',rent=12000)
        self.pay(ravi,'2024-01',4000)
        self.pay(ravi,'2024-01',6000)
        self.pay(ravi,'2024-07',5000)
        self.pay(priya,'2025-01',12000)
        self.assertEqual(main.dashboard('2024-01')['expected_rent'],10000)
        self.assertEqual(main.dashboard('2024-01')['pending_rent'],0)
        july = main.dashboard('2024-07')
        self.assertEqual(july['pending_rent'],6000)
        self.assertEqual(july['pending_properties'][0]['tenant_name'],'Ravi')
        self.assertEqual(main.dashboard('2025-01')['expected_rent'],12000)
        self.assertEqual(main.dashboard('2023-12')['expected_rent'],0)
        self.assertEqual({p['tenant_name'] for p in main.list_payments(rental_month='2024-01')},{'Ravi'})
        self.assertEqual(main.list_properties()[0]['tenant_name'],'Priya')
        api.update_rate(rate['id'],RentRateCreate(tenancy_id=ravi['id'],effective_month='2024-07',amount=11500))
        self.assertEqual(main.dashboard('2024-07')['pending_rent'],6500)
        self.assertEqual(main.dashboard('2024-01')['expected_rent'],10000)
        self.assertTrue(any(h['entity']=='rent_rates' for h in main.list_edit_history()))
        async def exported():
            return b''.join([part async for part in main.export_excel().body_iterator])
        workbook=load_workbook(io.BytesIO(asyncio.run(exported())))
        headers=[c.value for c in workbook['Rent History'][1]]
        tenants=[r[headers.index('Tenant')] for r in workbook['Rent History'].iter_rows(min_row=2,values_only=True)]
        self.assertIn('Ravi',tenants)
        self.assertIn('Priya',tenants)
        self.assertEqual(workbook['Rent Rates'].max_row,4)

    def test_overlaps_invalid_assignments_and_date_edits(self):
        stay=self.stay(self.ravi,'2024-01-15','2024-12-01')
        self.pay(stay,'2024-12',500)
        with self.assertRaises(HTTPException):
            self.stay(self.priya,'2024-12-02')
        for month in ['2023-12','2025-01','2024-1']:
            with self.assertRaises(HTTPException):
                self.pay(stay,month,100)
        with self.assertRaises(HTTPException):
            main.create_payment(PaymentCreate(property_id=self.prop['id'],rental_month='2024-01',amount=10,paid_on='2024-01-01'))
        with self.assertRaises(HTTPException):
            api.update_tenancy(stay['id'],TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id'],start_date='2024-01-01',end_date='2024-11-30',initial_rent=10000))
        self.assertEqual(api.list_tenancies()[0]['end_date'],'2024-12-01')
        other=main.create_property(PropertyCreate(name='Other',monthly_rent=100))
        with self.assertRaises(HTTPException):
            main.create_payment(PaymentCreate(property_id=other['id'],tenancy_id=stay['id'],rental_month='2024-01',amount=10,paid_on='2024-01-01'))
        api.create_rate(RentRateCreate(tenancy_id=stay['id'],effective_month='2024-07',amount=11000))
        with self.assertRaises(HTTPException):
            api.create_rate(RentRateCreate(tenancy_id=stay['id'],effective_month='2024-07',amount=12000))
        with self.assertRaises(HTTPException):
            api.create_rate(RentRateCreate(tenancy_id=stay['id'],effective_month='2025-01',amount=12000))

    def test_legacy_payments_explicit_assignment_and_no_cross_property_offset(self):
        pay=main.create_payment(PaymentCreate(property_id=self.prop['id'],rental_month='2024-01',amount=15000,paid_on='2024-02-01'))
        stay=self.stay(self.ravi,'2024-01-01','2024-12-31')
        self.assertEqual(main.dashboard('2024-01')['unassigned_payments'],1)
        self.assertEqual(main.dashboard('2024-01')['pending_rent'],10000)
        self.assertEqual(main.list_payments()[0]['tenant_name'],'Legacy person')
        main.edit_payment(pay['id'],PaymentCreate(property_id=self.prop['id'],tenancy_id=stay['id'],rental_month='2024-01',amount=15000,paid_on='2024-02-01'))
        self.assertEqual(main.dashboard('2024-01')['unassigned_payments'],0)
        other=main.create_property(PropertyCreate(name='Other',monthly_rent=1000))
        api.create_tenancy(TenancyCreate(property_id=other['id'],tenant_id=self.priya['id'],start_date='2024-01-01',initial_rent=1000))
        self.assertEqual(main.dashboard('2024-01')['pending_rent'],1000)
        database.init_db()
        self.assertEqual(len(api.list_tenancies()),2)
        self.assertEqual(main.list_payments()[0]['tenant_name'],'Ravi')
        self.assertTrue(any(c['field']=='tenancy_id' for h in main.list_edit_history() for c in h['changes']))

    def test_tenancy_edits_vacancies_and_inactive_historical_property(self):
        stay=self.stay(self.ravi,'2024-01-01','2024-12-31')
        api.update_tenant(self.ravi['id'],TenantCreate(name='Ravi corrected',phone='456'))
        api.update_tenancy(stay['id'],TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id'],start_date='2024-01-01',end_date='2024-12-31',initial_rent=9000,deposit=18000))
        main.update_property(self.prop['id'],PropertyUpdate(name='House',monthly_rent=999,active=False))
        self.assertEqual(main.dashboard('2024-05')['expected_rent'],9000)
        self.assertEqual(main.dashboard('2024-05')['pending_properties'][0]['tenant_name'],'Ravi corrected')
        self.assertEqual(main.dashboard('2025-01')['expected_rent'],0)
        self.assertEqual(main.list_properties()[0]['tenant_name'],'')
        self.assertEqual(main.list_properties()[0]['monthly_rent'],0)

    def test_upgrade_preserves_old_payments_without_inventing_tenancies(self):
        database.DB_PATH=Path(self.tmp.name)/'legacy.db'
        with sqlite3.connect(database.DB_PATH) as con:
            con.executescript("""
                CREATE TABLE subareas(id INTEGER PRIMARY KEY,name TEXT UNIQUE,city TEXT);
                INSERT INTO subareas VALUES(1,'Old area','City');
                CREATE TABLE properties(id INTEGER PRIMARY KEY,subarea_id INTEGER,name TEXT,address TEXT DEFAULT '',
                    tenant_name TEXT DEFAULT '',tenant_phone TEXT DEFAULT '',monthly_rent REAL,advance_paid REAL DEFAULT 0,
                    due_day INTEGER DEFAULT 5,active INTEGER DEFAULT 1);
                INSERT INTO properties(id,subarea_id,name,tenant_name,monthly_rent) VALUES(1,1,'Old house','Old tenant',1000);
                CREATE TABLE rent_payments(id INTEGER PRIMARY KEY,property_id INTEGER,rental_month TEXT,amount REAL,
                    paid_on TEXT,payment_method TEXT,reference TEXT,notes TEXT);
                INSERT INTO rent_payments VALUES(1,1,'2020-01',750,'2020-02-01','Cash','','Old record');
            """)
        database.init_db()
        database.init_db()
        payment=main.list_payments()[0]
        self.assertEqual(payment['amount'],750)
        self.assertEqual(payment['tenant_name'],'Old tenant')
        self.assertIsNone(payment['tenancy_id'])
        self.assertEqual(api.list_tenancies(),[])
        main.update_property(1,PropertyUpdate(name='Old house',tenant_name='Different name',monthly_rent=2000))
        self.assertEqual(main.list_payments()[0]['tenant_name'],'Old tenant')

    def test_unknown_dates_and_rent_can_be_completed_incrementally(self):
        stay=api.create_tenancy(TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id']))
        self.assertIsNone(stay['start_date'])
        self.assertIsNone(api.list_tenancies()[0]['initial_rent'])
        self.assertEqual(api.list_rates(),[])
        report=main.dashboard('2023-01')
        self.assertEqual(len(report['unknown_properties']),1)
        self.assertEqual(report['expected_rent'],0)
        self.assertIsNone(main.list_properties()[0]['monthly_rent'])
        self.pay(stay,'2023-07',4000)
        report=main.dashboard('2023-07')
        self.assertEqual(report['collected_rent'],4000)
        self.assertEqual(len(report['unknown_properties']),1)
        rate=api.create_rate(RentRateCreate(tenancy_id=stay['id'],effective_month='2023-07',amount=10000))
        self.assertEqual(main.dashboard('2023-07')['pending_rent'],6000)
        self.assertEqual(len(main.dashboard('2023-06')['unknown_properties']),1)
        api.create_rate(RentRateCreate(tenancy_id=stay['id'],effective_month='2023-01',amount=9000))
        api.update_tenancy(stay['id'],TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id'],start_date='2022-06-01'))
        self.assertEqual(main.dashboard('2023-01')['expected_rent'],9000)
        self.assertEqual(main.dashboard('2023-07')['expected_rent'],10000)
        self.assertEqual(len(main.dashboard('2022-06')['unknown_properties']),1)
        self.assertEqual(main.dashboard('2022-05')['unknown_properties'],[])
        api.update_tenancy(stay['id'],TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id'],start_date='2022-06-01',initial_rent=8000))
        self.assertEqual(main.dashboard('2022-06')['expected_rent'],8000)
        self.assertEqual(main.dashboard('2023-07')['expected_rent'],10000)
        database.init_db()
        self.assertEqual(len(api.list_rates()),3)

    def test_unknown_tenancies_do_not_double_charge_or_invent_occupancy(self):
        unknown=api.create_tenancy(TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id'],end_date='2024-12-31'))
        self.stay(self.priya,'2025-01-01',rent=12000)
        api.create_rate(RentRateCreate(tenancy_id=unknown['id'],effective_month='2024-01',amount=10000))
        self.assertEqual(main.dashboard('2025-01')['expected_rent'],12000)
        self.assertEqual(main.dashboard('2025-01')['unknown_properties'],[])
        self.assertEqual(main.dashboard('2024-01')['expected_rent'],10000)
        with self.assertRaises(HTTPException):
            self.pay(unknown,'2025-01',100)
        api.update_tenancy(unknown['id'],TenancyCreate(property_id=self.prop['id'],tenant_id=self.ravi['id']))
        report=main.dashboard('2025-01')
        self.assertEqual(report['expected_rent'],0)
        self.assertEqual(len(report['unknown_properties']),1)
        self.assertIn('More than one',report['unknown_properties'][0]['reason'])

    def test_delete_property_with_tenancy_is_blocked(self):
        stay=self.stay(self.ravi,'2024-01-01','2024-12-31')
        with self.assertRaises(HTTPException) as error:
            main.delete_property(self.prop['id'])
        self.assertEqual(error.exception.status_code,409)
        self.assertEqual(api.list_tenancies()[0]['id'],stay['id'])
        self.assertEqual(len(api.list_rates()),1)
        self.assertEqual(len(main.list_properties()),1)

    def test_deposit_refund_lifecycle_and_multi_installments(self):
        stay = self.stay(self.ravi, '2024-01-01', '2024-12-31') # deposit: 20000
        # Installment 1: 10000 refund
        ref1 = api.create_deposit_refund(DepositRefundCreate(
            tenancy_id=stay['id'],
            amount=10000,
            refunded_on='2025-01-05',
            payment_method='UPI',
            reference='UPI-12345'
        ))
        # Installment 2: 7000 refund with 3000 painting deduction
        ref2 = api.create_deposit_refund(DepositRefundCreate(
            tenancy_id=stay['id'],
            amount=7000,
            refunded_on='2025-01-20',
            payment_method='Bank transfer',
            deduction_amount=3000,
            deduction_reason='Painting & repair charges',
            notes='Final settlement'
        ))

        tenancies = api.list_tenancies()
        t = tenancies[0]
        self.assertEqual(t['deposit'], 20000)
        self.assertEqual(t['total_refunded'], 17000)
        self.assertEqual(t['total_deductions'], 3000)
        self.assertEqual(t['deposit_balance'], 0)
        self.assertTrue(t['is_settled'])
        self.assertEqual(len(t['refunds']), 2)

        # Audit history tracking
        api.update_deposit_refund(ref1['id'], DepositRefundCreate(
            tenancy_id=stay['id'],
            amount=11000,
            refunded_on='2025-01-05'
        ))
        self.assertTrue(any(h['entity'] == 'deposit_refunds' for h in main.list_edit_history()))

        # Delete refund
        api.delete_deposit_refund(ref2['id'])
        tenancies = api.list_tenancies()
        self.assertEqual(tenancies[0]['total_refunded'], 11000)
        self.assertEqual(tenancies[0]['total_deductions'], 0)
        self.assertEqual(tenancies[0]['deposit_balance'], 9000)
        self.assertFalse(tenancies[0]['is_settled'])

        # Test pure deduction (amount=0) with final settlement complete justified for whitewashing
        api.create_deposit_refund(DepositRefundCreate(
            tenancy_id=stay['id'],
            amount=0,
            refunded_on='2025-01-25',
            payment_method='No cash refund (Deductions only)',
            deduction_amount=9000,
            deduction_reason='Whitewashing & painting charges',
            is_final_settlement=True,
            notes='Full remaining deposit used for painting'
        ))
        tenancies = api.list_tenancies()
        self.assertEqual(tenancies[0]['deposit_balance'], 0)
        self.assertTrue(tenancies[0]['is_settled'])
        self.assertEqual(tenancies[0]['total_deductions'], 9000)

    def test_property_type_and_shop_business_name(self):
        shop = main.create_property(PropertyCreate(name='Shop 101', property_type='Shop', monthly_rent=25000))
        storage = main.create_property(PropertyCreate(name='Godown A', property_type='Storage unit', monthly_rent=8000))
        self.assertEqual(shop['property_type'], 'Shop')
        self.assertEqual(storage['property_type'], 'Storage unit')

        # Create shop tenancy with business name
        tenancy = api.create_tenancy(TenancyCreate(
            property_id=shop['id'],
            tenant_id=self.ravi['id'],
            business_name='Kathir Steels',
            start_date='2025-01-01',
            initial_rent=25000,
            deposit=50000
        ))
        self.assertEqual(tenancy['business_name'], 'Kathir Steels')

        # Check list_properties returns active business name and property type
        props = main.list_properties()
        shop_prop = next(p for p in props if p['id'] == shop['id'])
        self.assertEqual(shop_prop['property_type'], 'Shop')
        self.assertEqual(shop_prop['business_name'], 'Kathir Steels')
        self.assertEqual(shop_prop['tenant_name'], 'Ravi')

        # Check list_tenancies returns property_type and business_name
        tenancies = api.list_tenancies()
        shop_tenancy = next(t for t in tenancies if t['id'] == tenancy['id'])
        self.assertEqual(shop_tenancy['property_type'], 'Shop')
        self.assertEqual(shop_tenancy['business_name'], 'Kathir Steels')

        # Record payment and check list_payments
        pay = main.create_payment(PaymentCreate(
            property_id=shop['id'],
            tenancy_id=tenancy['id'],
            rental_month='2025-01',
            amount=25000,
            paid_on='2025-01-05'
        ))
        payments = main.list_payments(property_id=shop['id'])
        self.assertEqual(payments[0]['business_name'], 'Kathir Steels')
        self.assertEqual(payments[0]['property_type'], 'Shop')

        # Tenant rebrands / changes shop trade name
        api.update_tenancy(tenancy['id'], TenancyCreate(
            property_id=shop['id'],
            tenant_id=self.ravi['id'],
            business_name='Kathir Supermarket',
            start_date='2025-01-01',
            deposit=50000
        ))
        updated_tenancy = next(t for t in api.list_tenancies() if t['id'] == tenancy['id'])
        self.assertEqual(updated_tenancy['business_name'], 'Kathir Supermarket')
        self.assertTrue(any(h['entity'] == 'tenancies' and any(c['field'] == 'business_name' for c in h['changes']) for h in main.list_edit_history()))

        # Update property type
        main.update_property(shop['id'], PropertyUpdate(name='Shop 101', property_type='Commercial', monthly_rent=25000))
        self.assertTrue(any(h['entity'] == 'properties' and any(c['field'] == 'property_type' for c in h['changes']) for h in main.list_edit_history()))

        # Verify Excel export contains Property_Type and Business_Name columns
        async def exported():
            return b''.join([part async for part in main.export_excel().body_iterator])
        workbook = load_workbook(io.BytesIO(asyncio.run(exported())))
        prop_headers = [c.value for c in workbook['Properties'][1]]
        self.assertIn('Property_Type', prop_headers)
        self.assertIn('Business_Name', prop_headers)
        tenancy_headers = [c.value for c in workbook['Tenancies'][1]]
        self.assertIn('property_type', tenancy_headers)
        self.assertIn('business_name', tenancy_headers)


if __name__ == '__main__':
    unittest.main()


