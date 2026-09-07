import io
import json
import sqlite3
from datetime import date
from typing import Optional

from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import db, init_db
from .schemas import (PaymentCreate, PropertyCreate, PropertyUpdate,
                      GroupCreate, TaxCreate, ElectricityCreate, ElectricityCollection,
                      ExpenseCreate, ExpenseUpdate)

from .tenancies import router as tenancy_router, monthly_dashboard, validate_payment, list_tenants, list_tenancies, list_rates, list_deposit_refunds

app = FastAPI(title="Rent Ledger API", version="1.0.0")
from .archive import router as archive_router, list_archive, list_tracking, list_reviews
app.include_router(tenancy_router)
app.include_router(archive_router)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:4200"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup() -> None:
    init_db()


def rows(result):
    return [dict(row) for row in result.fetchall()]


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/dashboard")
def dashboard(month: Optional[str] = None):
    return monthly_dashboard(month)


@app.get("/api/properties")
def list_properties(group_id: Optional[int] = None, active: Optional[bool] = None):
    filters, params = [], []
    if group_id is not None:
        filters.append("p.group_id = ?")
        params.append(group_id)
    if active is not None:
        filters.append("p.active = ?")
        params.append(int(active))
    where = " WHERE " + " AND ".join(filters) if filters else ""
    with db() as connection:
        result = rows(connection.execute(
            """SELECT p.id, p.name, p.address, p.property_type, p.tenant_name, p.tenant_phone,
                      p.monthly_rent, p.advance_paid, p.due_day, p.active, p.group_id,
                      p.electricity_payer, g.name AS group_name,
                      (SELECT MAX(paid_on) FROM rent_payments rp WHERE rp.property_id = p.id) AS last_paid_on
               FROM properties p LEFT JOIN property_groups g ON g.id = p.group_id""" + where + " ORDER BY g.name, p.name",
            params,
        ))
        today = date.today().isoformat()
        for prop in result:
            prop['has_tenancy_history'] = bool(connection.execute('SELECT 1 FROM tenancies WHERE property_id=?',(prop['id'],)).fetchone())
            prop['business_name'] = ''
            if prop['has_tenancy_history']:
                candidates = connection.execute("""SELECT n.name,n.phone,t.deposit,t.business_name,
                    (SELECT amount FROM rent_rates WHERE tenancy_id=t.id AND effective_month<=? ORDER BY effective_month DESC LIMIT 1) AS rent
                    FROM tenancies t JOIN tenants n ON n.id=t.tenant_id WHERE t.property_id=? AND t.start_date<=?
                    AND (t.end_date IS NULL OR t.end_date>=?)""",(today[:7],prop['id'],today,today)).fetchall()
                current = candidates[0] if len(candidates)==1 else None
                prop.update(tenant_name=current['name'] if current else ('Tenancy unclear' if candidates else ''),
                            tenant_phone=current['phone'] if current else '',
                            business_name=current['business_name'] if current else '',
                            monthly_rent=current['rent'] if current else (None if candidates else 0),
                            advance_paid=current['deposit'] if current else 0)
        return result



@app.post("/api/properties", status_code=201)
def create_property(payload: PropertyCreate):
    data = payload.model_dump()
    data["active"] = int(data["active"])
    with db() as connection:
        if payload.group_id is not None and not connection.execute("SELECT 1 FROM property_groups WHERE id=?", (payload.group_id,)).fetchone():
            raise HTTPException(400, "Group not found")
        connection.execute("INSERT OR IGNORE INTO subareas(name, city) VALUES ('__legacy_storage__', '')")
        data['subarea_id'] = connection.execute("SELECT id FROM subareas WHERE name='__legacy_storage__'").fetchone()[0]
        try:
            cursor = connection.execute(
                """INSERT INTO properties(subarea_id, name, address, property_type, tenant_name, tenant_phone,
                   monthly_rent, advance_paid, due_day, active, group_id, electricity_payer) VALUES (:subarea_id, :name, :address,
                   :property_type, :tenant_name, :tenant_phone, :monthly_rent, :advance_paid, :due_day, :active, :group_id, :electricity_payer)""", data
            )
        except sqlite3.IntegrityError as error:
            raise HTTPException(400, str(error))
        return {"id": cursor.lastrowid, **payload.model_dump()}


@app.put("/api/properties/{property_id}")
def update_property(property_id: int, payload: PropertyUpdate):
    data = payload.model_dump() | {"id": property_id}
    data["active"] = int(data["active"])
    with db() as connection:
        if payload.group_id is not None and not connection.execute("SELECT 1 FROM property_groups WHERE id=?", (payload.group_id,)).fetchone():
            raise HTTPException(400, "Group not found")
        cursor = connection.execute(
            """UPDATE properties SET name=:name, address=:address, property_type=:property_type,
               tenant_name=:tenant_name, tenant_phone=:tenant_phone, monthly_rent=:monthly_rent,
               advance_paid=:advance_paid, due_day=:due_day, active=:active,
               group_id=:group_id, electricity_payer=:electricity_payer WHERE id=:id""", data
        )
        if not cursor.rowcount:
            raise HTTPException(404, "Property not found")
    return {"id": property_id, **payload.model_dump()}


@app.get("/api/payments")
def list_payments(property_id: Optional[int] = None, rental_month: Optional[str] = None):
    filters, params = [], []
    if property_id is not None:
        filters.append("rp.property_id = ?")
        params.append(property_id)
    if rental_month:
        filters.append("rp.rental_month = ?")
        params.append(rental_month)
    where = " WHERE " + " AND ".join(filters) if filters else ""
    with db() as connection:
        return rows(connection.execute(
            """SELECT rp.*, p.name AS property_name, p.property_type, COALESCE(n.name,rp.tenant_snapshot) AS tenant_name, t.business_name, g.name AS group_name
               FROM rent_payments rp JOIN properties p ON p.id = rp.property_id
               LEFT JOIN tenancies t ON t.id=rp.tenancy_id LEFT JOIN tenants n ON n.id=t.tenant_id
               LEFT JOIN property_groups g ON g.id = p.group_id""" + where + " ORDER BY rp.paid_on DESC, rp.id DESC",
            params,
        ))


@app.post("/api/payments", status_code=201)
def create_payment(payload: PaymentCreate):
    data = payload.model_dump(mode="json")
    with db() as connection:
        connection.execute("BEGIN IMMEDIATE")
        if not connection.execute("SELECT 1 FROM properties WHERE id = ?", (payload.property_id,)).fetchone():
            raise HTTPException(404, "Property not found")
        validate_payment(connection, data)
        data['paid_on'] = data['paid_on'] or ''
        data['tenant_snapshot'] = connection.execute("SELECT tenant_name FROM properties WHERE id=?", (payload.property_id,)).fetchone()[0]
        cursor = connection.execute(
            """INSERT INTO rent_payments(property_id, rental_month, amount, paid_on,
               payment_method, reference, notes, tenancy_id, tenant_snapshot) VALUES (:property_id, :rental_month, :amount,
               :paid_on, :payment_method, :reference, :notes, :tenancy_id, :tenant_snapshot)""", data
        )
        return {"id": cursor.lastrowid, **data}


@app.delete("/api/payments/{payment_id}", status_code=204)
def delete_payment(payment_id: int):
    with db() as connection:
        if not connection.execute("DELETE FROM rent_payments WHERE id = ?", (payment_id,)).rowcount:
            raise HTTPException(404, "Payment not found")


def style_sheet(sheet):
    fill = PatternFill("solid", fgColor="173F35")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


@app.get("/api/export")
def export_excel():
    with db() as connection:
        payments = rows(connection.execute(
            """SELECT rp.id AS Receipt_No, p.name AS Property, p.property_type AS Property_Type, g.name AS Property_Group, COALESCE(n.name,rp.tenant_snapshot) AS Tenant, t.business_name AS Business_Name, rp.tenancy_id AS Tenancy,
                      rp.rental_month AS Rent_Month, rp.amount AS Amount, rp.paid_on AS Paid_On,
                      rp.payment_method AS Method, rp.reference AS Reference, rp.notes AS Notes
               FROM rent_payments rp JOIN properties p ON p.id = rp.property_id
               LEFT JOIN tenancies t ON t.id=rp.tenancy_id LEFT JOIN tenants n ON n.id=t.tenant_id
               LEFT JOIN property_groups g ON g.id = p.group_id ORDER BY rp.paid_on, rp.id"""
        ))

    current_properties = list_properties()
    properties = [dict(Property=p['name'], Property_Type=p.get('property_type', 'House'), Address=p['address'],
                       Tenant=p['tenant_name'], Business_Name=p.get('business_name', ''), Phone=p['tenant_phone'],
                       Monthly_Rent=p['monthly_rent'], Advance_Paid=p['advance_paid'], Due_Day=p['due_day'],
                       Property_Group=p['group_name'], Electricity_Payer=p['electricity_payer'],
                       Status='Active' if p['active'] else 'Inactive') for p in current_properties]
    workbook = Workbook()
    prop_sheet = workbook.active
    prop_sheet.title = "Properties"
    prop_headers = list(properties[0]) if properties else ["Property", "Property_Type", "Address", "Tenant", "Business_Name", "Phone", "Monthly_Rent", "Advance_Paid", "Due_Day", "Property_Group", "Electricity_Payer", "Status"]
    prop_sheet.append(prop_headers)
    for item in properties:
        prop_sheet.append(list(item.values()))
    style_sheet(prop_sheet)

    pay_sheet = workbook.create_sheet("Rent History")
    pay_headers = list(payments[0]) if payments else ["Receipt_No", "Property", "Property_Type", "Property_Group", "Tenant", "Business_Name", "Tenancy", "Rent_Month", "Amount", "Paid_On", "Method", "Reference", "Notes"]
    pay_sheet.append(pay_headers)
    for item in payments:
        pay_sheet.append(list(item.values()))
    style_sheet(pay_sheet)

    for title, records, headers in [
        ("Archive", list_archive(), ["id","property_name","rental_month","tenant_label","expected_rent","amount","paid_on","source","notes","review_status","payment_id"]),
        ("Tracking", list_tracking(), ["id","start_month","opening_balance","tenant_label","notes"]),
        ("Monthly Review", list_reviews(), ["property_id","month","occupancy","review_status","expected_rent","notes"]),
        ("Tenants", list_tenants(), ["id","name","phone","notes"]),
        ("Tenancies", list_tenancies(), ["id","property_id","property_name","property_type","tenant_id","tenant_name","business_name","start_date","end_date","deposit","initial_rent","notes"]),
        ("Rent Rates", list_rates(), ["id","tenancy_id","effective_month","amount"]),
        ("Deposit Refunds", list_deposit_refunds(), ["id", "property_name", "tenant_name", "amount", "refunded_on", "payment_method", "reference", "deduction_amount", "deduction_reason", "is_final_settlement", "notes"]),
        ("Groups", list_groups(), ["id", "name", "kind", "property_count", "property_tax", "water_tax"]),
        ("Group Taxes", list_taxes(), ["id", "group_id", "tax_type", "amount", "paid_on", "period", "reference", "notes", "group_name"]),
        ("Electricity", list_electricity(), ["id", "property_id", "period", "amount", "paid_on", "collected", "reference", "notes", "property_name", "group_name"]),
        ("Expenses", list_expenses(), ["id", "property_id", "property_name", "group_name", "category", "amount", "expense_date", "paid_to", "payment_method", "reference", "notes"]),
    ]:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for record in records:
            sheet.append([record.get(key) for key in headers])
        style_sheet(sheet)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"rent-history-{date.today().isoformat()}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/groups")
def list_groups():
    with db() as connection:
        return rows(connection.execute("""SELECT g.*,
            (SELECT COUNT(*) FROM properties WHERE group_id=g.id) AS property_count,
            COALESCE((SELECT SUM(amount) FROM group_taxes WHERE group_id=g.id AND tax_type='Property tax'),0) AS property_tax,
            COALESCE((SELECT SUM(amount) FROM group_taxes WHERE group_id=g.id AND tax_type='Water tax'),0) AS water_tax
            FROM property_groups g ORDER BY g.name"""))


@app.post("/api/groups", status_code=201)
def create_group(payload: GroupCreate):
    with db() as connection:
        try:
            cursor = connection.execute("INSERT INTO property_groups(name,kind) VALUES (?,?)", (payload.name, payload.kind))
        except sqlite3.IntegrityError:
            raise HTTPException(409, "A group with this name already exists")
        return {"id": cursor.lastrowid, **payload.model_dump()}


@app.get("/api/taxes")
def list_taxes():
    with db() as connection:
        return rows(connection.execute("SELECT t.*, g.name AS group_name FROM group_taxes t JOIN property_groups g ON g.id=t.group_id ORDER BY t.paid_on DESC, t.id DESC"))


@app.post("/api/taxes", status_code=201)
def create_tax(payload: TaxCreate):
    with db() as connection:
        if not connection.execute("SELECT 1 FROM property_groups WHERE id=?", (payload.group_id,)).fetchone():
            raise HTTPException(404, "Group not found")
        data = payload.model_dump(mode="json")
        cursor = connection.execute("""INSERT INTO group_taxes(group_id,tax_type,amount,paid_on,period,reference,notes)
            VALUES (:group_id,:tax_type,:amount,:paid_on,:period,:reference,:notes)""", data)
        return {"id": cursor.lastrowid, **data}


@app.get("/api/electricity")
def list_electricity():
    with db() as connection:
        return rows(connection.execute("""SELECT b.*, p.name AS property_name,
            (SELECT name FROM property_groups WHERE id=p.group_id) AS group_name
            FROM electricity_bills b JOIN properties p ON p.id=b.property_id ORDER BY b.paid_on DESC,b.id DESC"""))


@app.post("/api/electricity", status_code=201)
def create_electricity(payload: ElectricityCreate):
    if payload.collected > payload.amount:
        raise HTTPException(400, "Collected amount cannot exceed the bill")
    with db() as connection:
        prop = connection.execute("SELECT electricity_payer FROM properties WHERE id=?", (payload.property_id,)).fetchone()
        if not prop:
            raise HTTPException(404, "Property not found")
        if prop['electricity_payer'] != 'Owner':
            raise HTTPException(400, "This property is set to tenant-paid electricity")
        data = payload.model_dump(mode="json")
        cursor = connection.execute("""INSERT INTO electricity_bills(property_id,period,amount,paid_on,collected,reference,notes)
            VALUES (:property_id,:period,:amount,:paid_on,:collected,:reference,:notes)""", data)
        return {"id": cursor.lastrowid, **data}


@app.put("/api/electricity/{bill_id}/collection")
def update_collection(bill_id: int, payload: ElectricityCollection):
    with db() as connection:
        bill = connection.execute("SELECT * FROM electricity_bills WHERE id=?", (bill_id,)).fetchone()
        if not bill:
            raise HTTPException(404, "Bill not found")
        if payload.collected > bill['amount']:
            raise HTTPException(400, "Collected amount cannot exceed the bill")
        connection.execute("UPDATE electricity_bills SET collected=? WHERE id=?", (payload.collected,bill_id))
        return {**dict(bill), "collected": payload.collected}


@app.delete("/api/taxes/{record_id}", status_code=204)
def delete_tax(record_id: int):
    with db() as connection:
        if not connection.execute("DELETE FROM group_taxes WHERE id=?", (record_id,)).rowcount:
            raise HTTPException(404, "Tax record not found")


@app.delete("/api/electricity/{record_id}", status_code=204)
def delete_electricity(record_id: int):
    with db() as connection:
        if not connection.execute("DELETE FROM electricity_bills WHERE id=?", (record_id,)).rowcount:
            raise HTTPException(404, "Bill not found")


@app.delete("/api/groups/{group_id}", status_code=204)
def delete_group(group_id: int):
    with db() as connection:
        connection.execute("BEGIN IMMEDIATE")
        if not connection.execute("SELECT 1 FROM property_groups WHERE id=?", (group_id,)).fetchone():
            raise HTTPException(404, "Group not found")
        if connection.execute("SELECT 1 FROM group_taxes WHERE group_id=?", (group_id,)).fetchone():
            raise HTTPException(409, "This group has tax records and cannot be deleted. Keep the group to preserve its tax history.")
        connection.execute("UPDATE properties SET group_id=NULL WHERE group_id=?", (group_id,))
        connection.execute("DELETE FROM property_groups WHERE id=?", (group_id,))


@app.get("/api/edit-history")
def list_edit_history():
    with db() as connection:
        records = rows(connection.execute("SELECT * FROM edit_history ORDER BY id DESC"))
    for record in records:
        before = json.loads(record.pop('before_json'))
        after = json.loads(record.pop('after_json'))
        record['label'] = after.get('name') or before.get('name') or f"#{record['record_id']}"
        record['changes'] = [dict(field=key, before=before[key], after=after[key])
                             for key in before if before[key] != after[key]]
    return records


def edit_record(table, record_id, payload):
    # Only server-owned table names are passed by the routes below.
    data = payload.model_dump(mode="json")
    with db() as connection:
        connection.execute("BEGIN IMMEDIATE")
        existing = connection.execute(f"SELECT * FROM {table} WHERE id=?", (record_id,)).fetchone()
        if not existing:
            raise HTTPException(404, "Record not found")
        if 'property_id' in data:
            prop = connection.execute("SELECT * FROM properties WHERE id=?", (data['property_id'],)).fetchone()
            if not prop:
                raise HTTPException(400, "Property not found")
            if table == 'electricity_bills':
                if data['collected'] > data['amount']:
                    raise HTTPException(400, "Collected amount cannot exceed the bill")
                if data['property_id'] != existing['property_id'] and prop['electricity_payer'] != 'Owner':
                    raise HTTPException(400, "Choose an owner-paid electricity property")
        if 'group_id' in data and not connection.execute("SELECT 1 FROM property_groups WHERE id=?", (data['group_id'],)).fetchone():
            raise HTTPException(400, "Group not found")
        if table == 'rent_payments':
            validate_payment(connection, data)
            data['paid_on'] = data['paid_on'] or ''
        assignments = ','.join(f'{field}=?' for field in data)
        try:
            connection.execute(f"UPDATE {table} SET {assignments} WHERE id=?", (*data.values(), record_id))
        except sqlite3.IntegrityError:
            raise HTTPException(409, "A group with this name already exists")
    return {"id": record_id, **data}


@app.put("/api/groups/{record_id}")
def edit_group(record_id: int, payload: GroupCreate):
    return edit_record('property_groups', record_id, payload)


@app.put("/api/payments/{record_id}")
def edit_payment(record_id: int, payload: PaymentCreate):
    return edit_record('rent_payments', record_id, payload)


@app.put("/api/taxes/{record_id}")
def edit_tax(record_id: int, payload: TaxCreate):
    return edit_record('group_taxes', record_id, payload)


@app.put("/api/electricity/{record_id}")
def edit_electricity(record_id: int, payload: ElectricityCreate):
    return edit_record('electricity_bills', record_id, payload)


def list_expenses(property_id: Optional[int] = None, group_id: Optional[int] = None, category: Optional[str] = None, from_date: Optional[str] = None, to_date: Optional[str] = None):
    filters, params = [], []
    if property_id is not None:
        filters.append("e.property_id = ?")
        params.append(property_id)
    if group_id is not None:
        filters.append("p.group_id = ?")
        params.append(group_id)
    if category:
        filters.append("e.category = ?")
        params.append(category)
    if from_date:
        filters.append("e.expense_date >= ?")
        params.append(from_date)
    if to_date:
        filters.append("e.expense_date <= ?")
        params.append(to_date)
    where = " WHERE " + " AND ".join(filters) if filters else ""
    with db() as connection:
        return rows(connection.execute(
            """SELECT e.*, p.name AS property_name, p.property_type, g.name AS group_name, g.id AS group_id
               FROM property_expenses e
               JOIN properties p ON p.id = e.property_id
               LEFT JOIN property_groups g ON g.id = p.group_id""" + where + " ORDER BY e.expense_date DESC, e.id DESC",
            params,
        ))


@app.get("/api/expenses")
def get_expenses(property_id: Optional[int] = None, group_id: Optional[int] = None, category: Optional[str] = None, from_date: Optional[str] = None, to_date: Optional[str] = None):
    return list_expenses(property_id, group_id, category, from_date, to_date)


@app.post("/api/expenses", status_code=201)
def create_expense(payload: ExpenseCreate):
    data = payload.model_dump(mode="json")
    with db() as connection:
        if not connection.execute("SELECT 1 FROM properties WHERE id = ?", (payload.property_id,)).fetchone():
            raise HTTPException(404, "Property not found")
        data['expense_date'] = data['expense_date'] or date.today().isoformat()
        cursor = connection.execute(
            """INSERT INTO property_expenses(property_id, category, amount, expense_date,
               paid_to, payment_method, reference, notes)
               VALUES (:property_id, :category, :amount, :expense_date, :paid_to, :payment_method, :reference, :notes)""",
            data
        )
        return {"id": cursor.lastrowid, **data}


@app.put("/api/expenses/{record_id}")
def edit_expense(record_id: int, payload: ExpenseUpdate):
    return edit_record('property_expenses', record_id, payload)


@app.delete("/api/expenses/{record_id}", status_code=204)
def delete_expense(record_id: int):
    with db() as connection:
        if not connection.execute("DELETE FROM property_expenses WHERE id = ?", (record_id,)).rowcount:
            raise HTTPException(404, "Expense record not found")


@app.delete("/api/properties/{property_id}", status_code=204)
def delete_property(property_id: int):
    with db() as connection:
        connection.execute("BEGIN IMMEDIATE")
        if not connection.execute("SELECT 1 FROM properties WHERE id=?", (property_id,)).fetchone():
            raise HTTPException(404, "Property not found")
        if connection.execute('SELECT 1 FROM tracking_settings WHERE id=?',(property_id,)).fetchone():
            raise HTTPException(409, 'This property has a tracking baseline. Mark it inactive to preserve the records.')
        for table in ('rent_payments', 'electricity_bills', 'property_expenses', 'tenancies', 'archive_entries', 'month_reviews'):
            if connection.execute(f"SELECT 1 FROM {table} WHERE property_id=?", (property_id,)).fetchone():
                raise HTTPException(409, "This property has linked financial, tenancy, or archive records and cannot be deleted. Use Edit details to mark it inactive and preserve its records.")
        connection.execute("DELETE FROM properties WHERE id=?", (property_id,))

