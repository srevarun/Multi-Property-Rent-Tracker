import io
import json
from datetime import date
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..database import db
from .properties import list_properties
from .groups_taxes import list_groups, list_taxes
from .electricity import list_electricity
from .expenses import list_expenses
from ..tenancies import list_tenants, list_tenancies, list_rates, list_deposit_refunds
from ..archive import list_archive, list_tracking, list_reviews

router = APIRouter()


def rows(result):
    return [dict(row) for row in result.fetchall()]


@router.get("/api/edit-history")
def list_edit_history():
    with db() as connection:
        records = rows(connection.execute("SELECT * FROM edit_history ORDER BY id DESC"))
    for record in records:
        before = json.loads(record.pop('before_json'))
        after = json.loads(record.pop('after_json'))
        record['label'] = after.get('name') or before.get('name') or f"#{record['record_id']}"
        record['changes'] = [dict(field=key, before=before[key], after=after[key])
                             for key in before if before[key] != after[key]]
    return records


def style_sheet(sheet):
    fill = PatternFill("solid", fgColor="173F35")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


@router.get("/api/export")
def export_excel():
    with db() as connection:
        payments = rows(connection.execute(
            """SELECT rp.id AS Receipt_No, p.name AS Property, p.property_type AS Property_Type, g.name AS Property_Group, COALESCE(n.name,rp.tenant_snapshot) AS Tenant, t.business_name AS Business_Name, rp.tenancy_id AS Tenancy,
                      rp.rental_month AS Rent_Month, rp.amount AS Amount, rp.paid_on AS Paid_On,
                      rp.payment_method AS Method, rp.reference AS Reference, rp.notes AS Notes
               FROM rent_payments rp JOIN properties p ON p.id = rp.property_id
               LEFT JOIN tenancies t ON t.id=rp.tenancy_id LEFT JOIN tenants n ON n.id=t.tenant_id
               LEFT JOIN property_groups g ON g.id = p.group_id ORDER BY rp.paid_on, rp.id"""
        ))

    current_properties = list_properties()
    properties = [dict(Property=p['name'], Property_Type=p.get('property_type', 'House'), Address=p['address'],
                       Tenant=p['tenant_name'], Business_Name=p.get('business_name', ''), Phone=p['tenant_phone'],
                       Monthly_Rent=p['monthly_rent'], Advance_Paid=p['advance_paid'], Due_Day=p['due_day'],
                       Property_Group=p['group_name'], Electricity_Payer=p['electricity_payer'],
                       Status='Active' if p['active'] else 'Inactive') for p in current_properties]
    workbook = Workbook()
    prop_sheet = workbook.active
    prop_sheet.title = "Properties"
    prop_headers = list(properties[0]) if properties else ["Property", "Property_Type", "Address", "Tenant", "Business_Name", "Phone", "Monthly_Rent", "Advance_Paid", "Due_Day", "Property_Group", "Electricity_Payer", "Status"]
    prop_sheet.append(prop_headers)
    for item in properties:
        prop_sheet.append(list(item.values()))
    style_sheet(prop_sheet)

    pay_sheet = workbook.create_sheet("Rent History")
    pay_headers = list(payments[0]) if payments else ["Receipt_No", "Property", "Property_Type", "Property_Group", "Tenant", "Business_Name", "Tenancy", "Rent_Month", "Amount", "Paid_On", "Method", "Reference", "Notes"]
    pay_sheet.append(pay_headers)
    for item in payments:
        pay_sheet.append(list(item.values()))
    style_sheet(pay_sheet)

    for title, records, headers in [
        ("Archive", list_archive(), ["id", "property_name", "rental_month", "tenant_label", "expected_rent", "amount", "paid_on", "source", "notes", "review_status", "payment_id"]),
        ("Tracking", list_tracking(), ["id", "start_month", "opening_balance", "tenant_label", "notes"]),
        ("Monthly Review", list_reviews(), ["property_id", "month", "occupancy", "review_status", "expected_rent", "notes"]),
        ("Tenants", list_tenants(), ["id", "name", "phone", "notes"]),
        ("Tenancies", list_tenancies(), ["id", "property_id", "property_name", "property_type", "tenant_id", "tenant_name", "business_name", "start_date", "end_date", "deposit", "initial_rent", "notes"]),
        ("Rent Rates", list_rates(), ["id", "tenancy_id", "effective_month", "amount"]),
        ("Deposit Refunds", list_deposit_refunds(), ["id", "property_name", "tenant_name", "amount", "refunded_on", "payment_method", "reference", "deduction_amount", "deduction_reason", "is_final_settlement", "notes"]),
        ("Groups", list_groups(), ["id", "name", "kind", "property_count", "property_tax", "water_tax"]),
        ("Group Taxes", list_taxes(), ["id", "group_id", "tax_type", "amount", "paid_on", "period", "reference", "notes", "group_name"]),
        ("Electricity", list_electricity(), ["id", "property_id", "period", "amount", "paid_on", "collected", "reference", "notes", "property_name", "group_name"]),
        ("Expenses", list_expenses(), ["id", "property_id", "property_name", "group_name", "category", "amount", "expense_date", "paid_to", "payment_method", "reference", "notes"]),
    ]:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for record in records:
            sheet.append([record.get(key) for key in headers])
        style_sheet(sheet)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"rent-history-{date.today().isoformat()}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})
