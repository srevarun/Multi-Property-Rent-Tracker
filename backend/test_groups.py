import asyncio
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from fastapi import HTTPException
from app import database, main
from app.schemas import (PropertyCreate, PropertyUpdate, GroupCreate,
                         TaxCreate, ElectricityCreate, ElectricityCollection, PaymentCreate)


class GroupLedgerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.original = database.DB_PATH
        database.DB_PATH = Path(self.tmp.name) / 'ledger.db'
        # Existing schema and data must survive the migration.
        with sqlite3.connect(database.DB_PATH) as con:
            con.executescript("""CREATE TABLE subareas(id INTEGER PRIMARY KEY, name TEXT UNIQUE, city TEXT);
                INSERT INTO subareas VALUES(1,'Central','City');
                CREATE TABLE properties(id INTEGER PRIMARY KEY, subarea_id INTEGER, name TEXT,
                address TEXT DEFAULT '', tenant_name TEXT DEFAULT '', tenant_phone TEXT DEFAULT '',
                monthly_rent REAL, advance_paid REAL DEFAULT 0, due_day INTEGER DEFAULT 5,
                active INTEGER DEFAULT 1);
                INSERT INTO properties(id,subarea_id,name,monthly_rent) VALUES(1,1,'Existing',1000);""")
        database.init_db()
        database.init_db()

    def tearDown(self):
        database.DB_PATH = self.original
        self.tmp.cleanup()

    def test_group_tax_and_electricity_lifecycle(self):
        self.assertEqual(main.list_properties()[0]['electricity_payer'], 'Tenant')
        self.assertIsNone(main.list_properties()[0]['group_id'])
        self.assertNotIn('subarea_id', main.list_properties()[0])
        group = main.create_group(GroupCreate(name='Market complex', kind='Shopping complex'))
        prop = main.create_property(PropertyCreate(name='Shop 1', monthly_rent=5000,
                                                   group_id=group['id'], electricity_payer='Owner'))
        for kind, amount in [('Property tax', 2000), ('Water tax', 500)]:
            main.create_tax(TaxCreate(group_id=group['id'], tax_type=kind, amount=amount,
                                     paid_on='2026-09-01', period='2026 H2'))
        g = main.list_groups()[0]
        self.assertEqual((g['property_count'],g['property_tax'],g['water_tax']), (1,2000,500))
        before = main.dashboard('2026-09')
        bill = main.create_electricity(ElectricityCreate(property_id=prop['id'], period='Jul-Aug',
                                      amount=1500, collected=500, paid_on='2026-09-02'))
        self.assertEqual(main.list_electricity()[0]['amount'] - main.list_electricity()[0]['collected'],1000)
        main.update_collection(bill['id'], ElectricityCollection(collected=1500))
        self.assertEqual(main.list_electricity()[0]['collected'],1500)
        self.assertEqual(main.dashboard('2026-09'), before)
        with self.assertRaises(HTTPException):
            main.update_collection(bill['id'], ElectricityCollection(collected=1501))
        with self.assertRaises(HTTPException):
            main.create_electricity(ElectricityCreate(property_id=1, period='Aug',amount=20,paid_on='2026-09-01'))
        # Changing future payer settings preserves already recorded owner bills.
        main.update_property(prop['id'], PropertyUpdate(name='Shop 1',monthly_rent=5000,
                              group_id=group['id'],electricity_payer='Tenant'))
        self.assertEqual(len(main.list_electricity()),1)
        self.assertEqual(main.list_properties(group_id=group['id'])[0]['group_name'], 'Market complex')
        async def export_bytes():
            return b''.join([chunk async for chunk in main.export_excel().body_iterator])
        workbook = load_workbook(io.BytesIO(asyncio.run(export_bytes())))
        self.assertEqual(workbook.sheetnames,['Properties','Rent History','Archive','Tracking','Monthly Review','Tenants','Tenancies','Rent Rates','Deposit Refunds','Groups','Group Taxes','Electricity','Expenses'])
        self.assertEqual(workbook['Group Taxes'].max_row,3)
        main.delete_electricity(bill['id'])
        main.delete_tax(main.list_taxes()[0]['id'])
        self.assertEqual(main.list_electricity(),[])
        self.assertEqual(len(main.list_taxes()),1)

    def test_invalid_groups_and_rent_regression(self):
        with self.assertRaises(HTTPException):
            main.create_property(PropertyCreate(name='Bad',monthly_rent=10,group_id=999))
        with self.assertRaises(HTTPException):
            main.create_tax(TaxCreate(group_id=999,tax_type='Water tax',amount=10,paid_on='2026-09-01'))
        main.create_payment(PaymentCreate(property_id=1,rental_month='2026-09',amount=400,paid_on='2026-09-01'))
        self.assertEqual(main.dashboard('2026-09')['pending_rent'],600)
        self.assertEqual(main.list_properties()[0]['name'],'Existing')

    def test_delete_group_preserves_property_ledgers(self):
        group = main.create_group(GroupCreate(name='Temporary group'))
        prop = main.create_property(PropertyCreate(name='Unit', monthly_rent=1000,
            group_id=group['id'], electricity_payer='Owner'))
        main.create_payment(PaymentCreate(property_id=prop['id'], rental_month='2026-09', amount=400, paid_on='2026-09-01'))
        main.create_electricity(ElectricityCreate(property_id=prop['id'], period='Aug', amount=100, paid_on='2026-09-01'))
        main.delete_group(group['id'])
        self.assertEqual(main.list_groups(), [])
        self.assertTrue(all(p['group_id'] is None for p in main.list_properties()))
        self.assertEqual(len(main.list_payments()), 1)
        self.assertEqual(len(main.list_electricity()), 1)
        with self.assertRaises(HTTPException) as error:
            main.delete_group(group['id'])
        self.assertEqual(error.exception.status_code, 404)
        empty = main.create_group(GroupCreate(name='Empty'))
        main.delete_group(empty['id'])
        self.assertEqual(main.list_groups(), [])

    def test_delete_group_with_taxes_is_blocked_without_changes(self):
        group = main.create_group(GroupCreate(name='Tax group'))
        main.create_property(PropertyCreate(name='Unit', monthly_rent=1000, group_id=group['id']))
        main.create_tax(TaxCreate(group_id=group['id'], tax_type='Water tax', amount=100, paid_on='2026-09-01'))
        with self.assertRaises(HTTPException) as error:
            main.delete_group(group['id'])
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(len(main.list_properties(group_id=group['id'])), 1)
        self.assertEqual(len(main.list_taxes()), 1)
        self.assertEqual(len(main.list_groups()), 1)

    def test_edits_and_persistent_history(self):
        group = main.create_group(GroupCreate(name='Before'))
        other = main.create_group(GroupCreate(name='Other'))
        prop = main.create_property(PropertyCreate(name='House', monthly_rent=1000, group_id=group['id'], electricity_payer='Owner'))
        pay = main.create_payment(PaymentCreate(property_id=prop['id'], rental_month='2026-09', amount=200, paid_on='2026-09-01'))
        tax = main.create_tax(TaxCreate(group_id=group['id'], tax_type='Water tax', amount=100, paid_on='2026-09-01'))
        bill = main.create_electricity(ElectricityCreate(property_id=prop['id'], period='Aug', amount=100, paid_on='2026-09-01'))
        self.assertEqual(main.list_edit_history(), [])
        main.edit_group(group['id'], GroupCreate(name='After', kind='Complex'))
        updated = PropertyUpdate(name='House revised', monthly_rent=1200, group_id=group['id'], electricity_payer='Tenant')
        main.update_property(prop['id'], updated)
        main.edit_payment(pay['id'], PaymentCreate(property_id=prop['id'], rental_month='2026-08', amount=300, paid_on='2026-09-02'))
        main.edit_tax(tax['id'], TaxCreate(group_id=other['id'], tax_type='Property tax', amount=150, paid_on='2026-09-02'))
        # Historical electricity bills remain editable after changing the payer.
        main.edit_electricity(bill['id'], ElectricityCreate(property_id=prop['id'], period='Aug revised', amount=200, collected=50, paid_on='2026-09-02'))
        main.update_collection(bill['id'], ElectricityCollection(collected=100))
        history = main.list_edit_history()
        self.assertEqual(len(history), 6)
        self.assertEqual(history[0]['changes'], [{'field':'collected','before':50,'after':100}])
        self.assertTrue(history[0]['edited_at'].endswith('Z'))
        self.assertEqual(main.dashboard('2026-09')['collected_rent'], 0)
        self.assertEqual(main.dashboard('2026-08')['collected_rent'], 300)
        self.assertEqual(main.list_taxes()[0]['group_id'], other['id'])
        main.update_property(prop['id'], updated)
        self.assertEqual(len(main.list_edit_history()), 6)  # No-op saves are not edits.
        with self.assertRaises(HTTPException):
            main.edit_group(group['id'], GroupCreate(name='Other'))
        with self.assertRaises(HTTPException):
            main.edit_electricity(bill['id'], ElectricityCreate(property_id=prop['id'], period='Aug', amount=10, collected=11, paid_on='2026-09-01'))
        with self.assertRaises(HTTPException):
            main.edit_payment(pay['id'], PaymentCreate(property_id=999, rental_month='2026-09', amount=200, paid_on='2026-09-01'))
        with self.assertRaises(HTTPException):
            main.edit_tax(999, TaxCreate(group_id=group['id'], tax_type='Water tax', amount=100, paid_on='2026-09-01'))
        self.assertEqual(len(main.list_edit_history()), 6)
        self.assertEqual(main.list_electricity()[0]['amount'], 200)
        main.delete_payment(pay['id'])
        database.init_db()
        self.assertEqual(len(main.list_edit_history()), 6)  # Survives deletion and restart.

    def test_delete_empty_property_preserves_group_and_edit_history(self):
        group = main.create_group(GroupCreate(name='Apartment'))
        prop = main.create_property(PropertyCreate(name='Empty unit',monthly_rent=1000,group_id=group['id']))
        main.update_property(prop['id'],PropertyUpdate(name='Renamed unit',monthly_rent=1000,group_id=group['id']))
        history = main.list_edit_history()
        main.delete_property(prop['id'])
        self.assertFalse(any(p['id']==prop['id'] for p in main.list_properties()))
        self.assertEqual(main.list_groups()[0]['property_count'],0)
        self.assertEqual(main.list_edit_history(),history)
        with self.assertRaises(HTTPException) as error:
            main.delete_property(prop['id'])
        self.assertEqual(error.exception.status_code,404)

    def test_delete_property_with_financial_records_is_blocked(self):
        main.create_payment(PaymentCreate(property_id=1,rental_month='2026-09',amount=500,paid_on='2026-09-01'))
        with self.assertRaises(HTTPException) as error:
            main.delete_property(1)
        self.assertEqual(error.exception.status_code,409)
        self.assertEqual(len(main.list_payments()),1)
        prop=main.create_property(PropertyCreate(name='EB unit',monthly_rent=1000,electricity_payer='Owner'))
        main.create_electricity(ElectricityCreate(property_id=prop['id'],period='Aug',amount=100,paid_on='2026-09-01'))
        with self.assertRaises(HTTPException) as error:
            main.delete_property(prop['id'])
        self.assertEqual(error.exception.status_code,409)
        self.assertEqual(len(main.list_electricity()),1)
        self.assertEqual(len(main.list_properties()),2)


if __name__ == '__main__':
    unittest.main()
